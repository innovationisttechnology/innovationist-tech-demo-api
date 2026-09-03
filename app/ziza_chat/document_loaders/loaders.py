"""Per-file-type extraction into text and images.

Every loader returns the same LoadedDocument shape, so the ingestion pipeline
downstream is identical regardless of file type: text goes to the chunker,
images go to the vision agent to be captioned and then chunked as text.
"""

import io
import logging
import mimetypes
from pathlib import PurePosixPath

from app.ziza_chat.document_loaders.base import (
    SUPPORTED_IMAGE_TYPES,
    ExtractedImage,
    ExtractedText,
    LoadedDocument,
    UnsupportedDocumentError,
)

logger = logging.getLogger(__name__)

MIN_EMBEDDED_IMAGE_BYTES = 8_000

SCANNED_PAGE_TEXT_THRESHOLD = 40

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".csv", ".json", ".yaml", ".yml"}

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

# Every OOXML file is a ZIP, so the signature cannot tell .docx from .xlsx,
# .jar, or a plain .zip — puremagic reports all of them as DOCX. The extension
# is the only disambiguator for these.
ZIP_CONTAINER_TYPES = {
    DOCX_MEDIA_TYPE,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "application/zip",
}

SNIFFABLE_TEXT_TYPES = {"application/json", "text/plain", "text/csv"}

TEXT_SAMPLE_BYTES = 65_536
MAX_REPLACEMENT_RATIO = 0.01


def guess_media_type(filename: str, declared: str | None = None) -> str:
    """Name-based guess. Only a hint — see detect_media_type."""
    if declared and declared != "application/octet-stream":
        return declared
    guessed, _ = mimetypes.guess_type(filename)
    return guessed or "application/octet-stream"


def sniff_media_type(data: bytes) -> str | None:
    """Media type read from the file's own bytes, or None when it has no
    recognisable signature (true for plain text — and for random binary)."""
    import puremagic

    try:
        detected = puremagic.from_string(data, mime=True)
    except Exception:
        return None
    return detected or None


def looks_like_text(data: bytes) -> bool:
    sample = data[:TEXT_SAMPLE_BYTES]
    if not sample or b"\x00" in sample:
        return False
    decoded = sample.decode("utf-8", errors="replace")
    return decoded.count("�") / len(decoded) <= MAX_REPLACEMENT_RATIO


def detect_media_type(
    data: bytes, filename: str, declared: str | None = None
) -> str:
    """Decide the type from the file's *contents*, not its name.

    A filename and a client-supplied content type are both attacker-controlled,
    so an executable renamed to .txt would otherwise be handed to a loader.
    Contents win wherever they carry a signature.
    """
    named_type = guess_media_type(filename, declared)
    suffix = PurePosixPath(filename).suffix.lower()
    sniffed = sniff_media_type(data)

    if sniffed is None:
        if looks_like_text(data) and (
            suffix in TEXT_EXTENSIONS or named_type.startswith("text/")
        ):
            return named_type if named_type.startswith("text/") else "text/plain"
        raise UnsupportedDocumentError(
            f"Cannot extract {filename!r}: its contents match no supported "
            "format. Supported: plain text, Markdown, CSV/JSON/YAML, PDF, "
            "DOCX, PNG/JPEG/GIF/WebP."
        )

    if sniffed in ZIP_CONTAINER_TYPES:
        if suffix == ".docx":
            return DOCX_MEDIA_TYPE
        raise UnsupportedDocumentError(
            f"Cannot extract {filename!r}: it is a ZIP-based archive and only "
            ".docx is supported among those."
        )

    if sniffed in SUPPORTED_IMAGE_TYPES or sniffed == "application/pdf":
        if named_type != sniffed:
            logger.warning(
                "%s contains %s but its name suggests %s — trusting contents",
                filename,
                sniffed,
                named_type,
            )
        return sniffed

    if sniffed in SNIFFABLE_TEXT_TYPES and looks_like_text(data):
        return "text/plain"

    raise UnsupportedDocumentError(
        f"Cannot extract {filename!r} (contents detected as {sniffed}). "
        "Supported: plain text, Markdown, CSV/JSON/YAML, PDF, DOCX, "
        "PNG/JPEG/GIF/WebP."
    )


def load_document(
    data: bytes, filename: str, content_type: str | None = None
) -> LoadedDocument:
    media_type = detect_media_type(data, filename, content_type)

    if media_type in SUPPORTED_IMAGE_TYPES:
        return LoadedDocument(images=[ExtractedImage(data=data, media_type=media_type)])

    try:
        if media_type == "application/pdf":
            return load_pdf(data)
        if media_type == DOCX_MEDIA_TYPE:
            return load_docx(data)
        return load_plain_text(data)
    except UnsupportedDocumentError:
        raise
    except Exception as failure:
        # The parser's own message names the library that raised it, so it is
        # logged rather than returned.
        logger.warning("Failed to parse %s as %s: %s", filename, media_type, failure)
        raise UnsupportedDocumentError(
            f"Could not read {filename!r} as {media_type} — it may be damaged, "
            "password-protected, or not really that format."
        ) from failure


def load_plain_text(data: bytes) -> LoadedDocument:
    return LoadedDocument(texts=[ExtractedText(text=data.decode("utf-8", "replace"))])


def load_pdf(data: bytes) -> LoadedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    texts: list[ExtractedText] = []
    images: list[ExtractedImage] = []

    for page_index, page in enumerate(reader.pages, start=1):
        locator = f"page {page_index}"
        page_text = (page.extract_text() or "").strip()
        page_images = _extract_pdf_page_images(page, locator)

        if len(page_text) >= SCANNED_PAGE_TEXT_THRESHOLD:
            texts.append(ExtractedText(text=page_text, locator=locator))
            images.extend(page_images)
            continue

        if page_images:
            logger.info("PDF %s looks scanned; routing to vision", locator)
            images.extend(page_images)
        elif page_text:
            texts.append(ExtractedText(text=page_text, locator=locator))

    return LoadedDocument(texts=texts, images=images)


def _extract_pdf_page_images(page: object, locator: str) -> list[ExtractedImage]:
    extracted: list[ExtractedImage] = []
    try:
        page_images = list(page.images)  # type: ignore[attr-defined]
    except Exception as failure:  # pypdf raises broadly on damaged XObjects
        logger.warning("Could not read images on %s: %s", locator, failure)
        return extracted

    for image in page_images:
        image_data = getattr(image, "data", None)
        if not image_data or len(image_data) < MIN_EMBEDDED_IMAGE_BYTES:
            continue
        media_type = guess_media_type(getattr(image, "name", "") or "image.png")
        if media_type not in SUPPORTED_IMAGE_TYPES:
            continue
        extracted.append(
            ExtractedImage(data=image_data, media_type=media_type, locator=locator)
        )
    return extracted


def load_docx(data: bytes) -> LoadedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))

    blocks = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    texts = [ExtractedText(text="\n\n".join(blocks))] if blocks else []

    images: list[ExtractedImage] = []
    for part in document.part.related_parts.values():
        blob = getattr(part, "blob", None)
        content_type = getattr(part, "content_type", "")
        if (
            blob
            and content_type in SUPPORTED_IMAGE_TYPES
            and len(blob) >= MIN_EMBEDDED_IMAGE_BYTES
        ):
            images.append(ExtractedImage(data=blob, media_type=content_type))

    return LoadedDocument(texts=texts, images=images)
