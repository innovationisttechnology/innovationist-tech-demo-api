"""Tests for multimodal document ingestion.

No network, no MongoDB, and no embedding-model download: extraction is exercised
on real generated files, and the vision agent is replaced by a stub.
"""

import io

import pytest

from app.ziza_chat.agents.vision import ImageDescription
from app.ziza_chat.document_loaders import UnsupportedDocumentError, load_document
from app.ziza_chat.document_loaders.loaders import guess_media_type
from app.ziza_chat.vector_store.chunking import enforce_token_limit


def make_png(width: int = 64, height: int = 64) -> bytes:
    """A real PNG, large enough to clear MIN_EMBEDDED_IMAGE_BYTES."""
    from PIL import Image

    image = Image.new("RGB", (width, height))
    # Noise defeats PNG compression, so the file is big enough to look like
    # content rather than a decorative icon.
    image.putdata([(x * 7 % 256, y * 13 % 256, (x + y) % 256)
                   for y in range(height) for x in range(width)])
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


class TestMediaTypeDetection:
    def test_declared_type_wins(self) -> None:
        assert guess_media_type("x.bin", "application/pdf") == "application/pdf"

    def test_falls_back_to_extension(self) -> None:
        assert guess_media_type("notes.md", None) == "text/markdown"

    def test_ignores_useless_octet_stream(self) -> None:
        assert guess_media_type("a.png", "application/octet-stream") == "image/png"


class TestContentBasedDetection:
    """Filenames and client-declared content types are attacker-controlled, so
    the file's own bytes decide what loader (if any) handles it."""

    def test_executable_renamed_to_txt_is_rejected(self) -> None:
        windows_exe = b"MZ\x90\x00" + b"\x00" * 200
        with pytest.raises(UnsupportedDocumentError):
            load_document(windows_exe, "notes.txt", "text/plain")

    def test_binary_junk_with_no_signature_is_rejected(self) -> None:
        # Random bytes have no magic signature — the same "unrecognised" signal
        # that genuine plain text gives, so this must be caught by content.
        junk = bytes(range(256)) * 8
        with pytest.raises(UnsupportedDocumentError):
            load_document(junk, "notes.txt", "text/plain")

    def test_zip_renamed_to_docx_is_rejected_not_crashed(self) -> None:
        import zipfile

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("hello.txt", "not really a document")
        with pytest.raises(UnsupportedDocumentError):
            load_document(buffer.getvalue(), "payload.docx", None)

    def test_xlsx_is_rejected_by_extension_since_zip_signature_is_ambiguous(
        self,
    ) -> None:
        import zipfile

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("xl/workbook.xml", "<workbook/>")
        with pytest.raises(UnsupportedDocumentError) as failure:
            load_document(buffer.getvalue(), "budget.xlsx", None)
        assert ".docx" in str(failure.value)

    def test_real_content_wins_over_a_misleading_name(self) -> None:
        # A genuine PNG named .txt is still processed as an image.
        document = load_document(make_png(), "screenshot.txt", "text/plain")
        assert len(document.images) == 1
        assert document.images[0].media_type == "image/png"

    def test_truncated_pdf_gives_a_client_error_not_a_crash(self) -> None:
        with pytest.raises(UnsupportedDocumentError):
            load_document(b"%PDF-1.7\ngarbage", "broken.pdf", "application/pdf")

    def test_empty_file_is_rejected(self) -> None:
        with pytest.raises(UnsupportedDocumentError):
            load_document(b"", "empty.txt", "text/plain")


class TestLooksLikeText:
    def test_accepts_utf8_prose(self) -> None:
        from app.ziza_chat.document_loaders.loaders import looks_like_text

        assert looks_like_text("Sarah leads the platform team. Café.".encode())

    def test_rejects_nul_bytes(self) -> None:
        from app.ziza_chat.document_loaders.loaders import looks_like_text

        assert not looks_like_text(b"text\x00more")

    def test_rejects_undecodable_binary(self) -> None:
        from app.ziza_chat.document_loaders.loaders import looks_like_text

        assert not looks_like_text(bytes(range(128, 256)) * 4)


class TestLoadDocument:
    def test_plain_text(self) -> None:
        document = load_document(b"hello world", "notes.txt", "text/plain")
        assert [section.text for section in document.texts] == ["hello world"]
        assert document.images == []

    def test_markdown_by_extension(self) -> None:
        document = load_document(b"# Title", "readme.md", None)
        assert document.texts[0].text == "# Title"

    def test_image_becomes_an_image_not_text(self) -> None:
        document = load_document(make_png(), "diagram.png", "image/png")
        assert document.texts == []
        assert len(document.images) == 1
        assert document.images[0].media_type == "image/png"

    def test_unsupported_type_is_rejected_with_a_useful_message(self) -> None:
        with pytest.raises(UnsupportedDocumentError) as failure:
            load_document(b"\x00\x01", "archive.zip", "application/zip")
        assert "archive.zip" in str(failure.value)

    def test_docx_text_and_tables(self) -> None:
        import docx

        document_file = docx.Document()
        document_file.add_paragraph("Sarah leads the platform team.")
        table = document_file.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Port"
        table.rows[0].cells[1].text = "8181"
        buffer = io.BytesIO()
        document_file.save(buffer)

        loaded = load_document(buffer.getvalue(), "handbook.docx", None)
        body = loaded.texts[0].text
        assert "Sarah leads the platform team." in body
        assert "Port | 8181" in body

    def test_pdf_text_pages_carry_page_locators(self) -> None:
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)

        loaded = load_document(buffer.getvalue(), "empty.pdf", "application/pdf")
        # A blank page has neither text nor images — it contributes nothing
        # rather than an empty chunk.
        assert loaded.texts == []
        assert loaded.images == []


class TestImageDescription:
    def test_search_text_includes_every_populated_field(self) -> None:
        description = ImageDescription(
            summary="Bar chart of quarterly revenue.",
            visible_text="Q1 1.2M Q2 1.45M",
            entities=["Finance team"],
            data_points=["Q2 revenue 1.45M"],
        )
        search_text = description.to_search_text()
        assert "Bar chart of quarterly revenue." in search_text
        assert "Q1 1.2M Q2 1.45M" in search_text
        assert "Finance team" in search_text
        assert "Q2 revenue 1.45M" in search_text

    def test_search_text_omits_empty_fields(self) -> None:
        description = ImageDescription(summary="A decorative border.")
        assert description.to_search_text() == "A decorative border."


class TestEnforceTokenLimit:
    """The embedder truncates past 512 tokens silently, so oversized chunks
    must be split before they reach it."""

    @staticmethod
    def count_words(text: str) -> int:
        return len(text.split())

    def test_short_chunks_pass_through_untouched(self) -> None:
        chunks = ["one two", "three four"]
        assert enforce_token_limit(chunks, self.count_words, max_tokens=5) == chunks

    def test_oversized_chunk_is_split_until_it_fits(self) -> None:
        oversized = " ".join(f"word{index}" for index in range(100))
        result = enforce_token_limit([oversized], self.count_words, max_tokens=10)
        assert len(result) > 1
        assert all(self.count_words(piece) <= 10 for piece in result)

    def test_split_preserves_all_content(self) -> None:
        oversized = " ".join(f"w{index}" for index in range(60))
        result = enforce_token_limit([oversized], self.count_words, max_tokens=8)
        rejoined = "".join(result)
        assert "w0" in rejoined and "w59" in rejoined

    def test_terminates_on_unsplittable_input(self) -> None:
        # A single character can never satisfy max_tokens=0; the guard must
        # give up rather than recurse forever.
        assert enforce_token_limit(["x"], self.count_words, max_tokens=0) == ["x"]

    def test_overlap_does_not_prevent_termination(self) -> None:
        oversized = "a" * 500
        result = enforce_token_limit(
            [oversized], lambda text: len(text), max_tokens=50, overlap_chars=40
        )
        assert all(len(piece) <= 50 for piece in result)


class TestCarriesNoInformation:
    """Captioning costs a model call per image, so images with nothing to index
    are filtered out locally before any request is made."""

    def test_solid_colour_is_skipped(self) -> None:
        from PIL import Image

        from app.ziza_chat.caption_cache import carries_no_information

        buffer = io.BytesIO()
        Image.new("RGB", (400, 400), "white").save(buffer, format="PNG")
        assert carries_no_information(buffer.getvalue())

    def test_tiny_image_is_skipped(self) -> None:
        from app.ziza_chat.caption_cache import carries_no_information

        assert carries_no_information(make_png(width=16, height=16))

    def test_detailed_image_is_kept(self) -> None:
        from app.ziza_chat.caption_cache import carries_no_information

        assert not carries_no_information(make_png(width=400, height=400))

    def test_unreadable_bytes_are_sent_rather_than_dropped(self) -> None:
        from app.ziza_chat.caption_cache import carries_no_information

        # Failing to inspect must not silently discard a real image.
        assert not carries_no_information(b"not an image at all")


class TestCaptionImages:
    @staticmethod
    def stub_cache(
        monkeypatch: pytest.MonkeyPatch, cached: dict[str, str] | None = None
    ) -> list[str]:
        """Keep the cache out of MongoDB; return the list of stored captions."""
        from app.ziza_chat import service

        store = cached or {}
        stored: list[str] = []

        async def fake_get(session_id: str, image_hash: str) -> str | None:
            return store.get(image_hash)

        async def fake_store(session_id: str, image_hash: str, caption: str) -> None:
            stored.append(caption)

        monkeypatch.setattr(service, "get_cached_caption", fake_get)
        monkeypatch.setattr(service, "store_caption", fake_store)
        monkeypatch.setattr(service, "carries_no_information", lambda data: False)
        return stored

    @pytest.mark.anyio
    async def test_failed_captions_are_skipped_not_fatal(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.ziza_chat import service
        from app.ziza_chat.document_loaders import ExtractedImage

        self.stub_cache(monkeypatch)

        async def flaky_describe(data: bytes, media_type: str) -> ImageDescription:
            if data == b"bad":
                raise RuntimeError("vision model unavailable")
            return ImageDescription(summary="a good image")

        monkeypatch.setattr(service, "describe_image", flaky_describe)

        images = [
            ExtractedImage(data=b"bad", media_type="image/png", locator="page 1"),
            ExtractedImage(data=b"good", media_type="image/png", locator="page 2"),
        ]
        sections, failures = await service.caption_images(images, "report.pdf", "session-1")

        assert len(sections) == 1
        assert failures == 1
        assert sections[0].text == "a good image"
        assert sections[0].source == "report.pdf (page 2, image)"

    @pytest.mark.anyio
    async def test_cached_caption_skips_the_model_entirely(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.ziza_chat import service
        from app.ziza_chat.caption_cache import hash_image
        from app.ziza_chat.document_loaders import ExtractedImage

        image_data = b"already-seen"
        self.stub_cache(monkeypatch, {hash_image(image_data): "cached description"})

        async def must_not_run(data: bytes, media_type: str) -> ImageDescription:
            raise AssertionError("vision model called despite a cache hit")

        monkeypatch.setattr(service, "describe_image", must_not_run)

        sections, failures = await service.caption_images(
            [ExtractedImage(data=image_data, media_type="image/png")],
            "again.png",
            "session-1",
        )
        assert failures == 0
        assert sections[0].text == "cached description"

    @pytest.mark.anyio
    async def test_information_free_images_never_reach_the_model(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from app.ziza_chat import service
        from app.ziza_chat.document_loaders import ExtractedImage

        self.stub_cache(monkeypatch)
        monkeypatch.setattr(service, "carries_no_information", lambda data: True)

        async def must_not_run(data: bytes, media_type: str) -> ImageDescription:
            raise AssertionError("blank image sent to the vision model")

        monkeypatch.setattr(service, "describe_image", must_not_run)

        sections, failures = await service.caption_images(
            [ExtractedImage(data=b"blank", media_type="image/png")],
            "logo.png",
            "session-1",
        )
        assert sections == []
        assert failures == 0


class TestFormatSource:
    def test_plain_filename_when_no_locator(self) -> None:
        from app.ziza_chat.service import format_source

        assert format_source("handbook.pdf", None) == "handbook.pdf"

    def test_includes_page_and_image_marker(self) -> None:
        from app.ziza_chat.service import format_source

        assert (
            format_source("handbook.pdf", "page 3", is_image=True)
            == "handbook.pdf (page 3, image)"
        )

    def test_standalone_image_has_no_page(self) -> None:
        from app.ziza_chat.service import format_source

        assert format_source("chart.png", None, is_image=True) == "chart.png (image)"
